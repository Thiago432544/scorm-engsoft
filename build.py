"""Gera o pacote SCORM 1.2 a partir do banco de questões V/F exportado do Google Forms (.xlsx).

Uso: python build.py caminho/para/banco.xlsx

Atualiza pacote/questoes.js e gera dist/SCORM_EngSoft.zip pronto para subir no LMS.
Só as colunas da questão são lidas; nome, e-mail e matrícula de quem enviou ficam de fora.
"""
import json, os, re, sys, zipfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PACOTE = os.path.join(HERE, "pacote")
OUT_ZIP = os.path.join(HERE, "dist", "SCORM_EngSoft.zip")


def clean(s):
    if s is None:
        return ""
    s = str(s).replace("\r\n", "\n").strip()
    return re.sub(r"[ \t]+\n", "\n", s)


def ler_questoes(xlsx):
    ws = openpyxl.load_workbook(xlsx, read_only=True).active
    header = next(ws.iter_rows(max_row=1, values_only=True))
    col = {name: i for i, name in enumerate(header)}

    questoes, vistos = [], set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        texto = clean(row[col["Questão"]])
        resp = clean(row[col["Resposta"]]).lower()
        if not texto or resp not in ("verdadeira", "falsa"):
            continue
        chave = re.sub(r"\s+", " ", texto).lower()
        if chave in vistos:
            continue
        vistos.add(chave)
        questoes.append({
            "id": len(questoes) + 1,
            "topico": clean(row[col["Tópico da questão"]]),
            "texto": texto,
            "resposta": resp == "verdadeira",
            "justificativa": clean(row[col["Citações e referências"]]),
        })
    return questoes


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    questoes = ler_questoes(sys.argv[1])
    print(f"{len(questoes)} questões únicas")

    with open(os.path.join(PACOTE, "questoes.js"), "w", encoding="utf-8") as fh:
        fh.write("window.QUESTOES = ")
        json.dump(questoes, fh, ensure_ascii=False, indent=0)
        fh.write(";\n")

    os.makedirs(os.path.dirname(OUT_ZIP), exist_ok=True)
    with zipfile.ZipFile(OUT_ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(os.listdir(PACOTE)):
            z.write(os.path.join(PACOTE, f), f)
    print("ok:", OUT_ZIP)


if __name__ == "__main__":
    main()
